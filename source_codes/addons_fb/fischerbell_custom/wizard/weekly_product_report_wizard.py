import base64
import calendar

from odoo import _, api, models, fields
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Border, Side, Alignment, Font
from io import BytesIO
from datetime import datetime, timedelta


class WeeklyProductReportwizard(models.TransientModel):
    _name = "weekly.product.report.wizard"
    _description = "Weekly Product Report Wizard"

    years = {'2015': '2015'}

    def _get_years(self):
        last = max(self.years.keys())
        current = datetime.now().year
        if int(last) < current:
            self.years.update(dict([[str(i), str(i)] for i in range(int(last) + 1, current + 1)]))
        return [(year, year) for year in self.years]

    year = fields.Selection(_get_years, "Year", required=True)

    def action_create_weekly_report(self):
        wb = Workbook()
        f_name = 'Sales History'

        table_header_alignment = Alignment(horizontal='left', vertical='top', text_rotation=0, wrap_text=True)
        bold_text = Font(name="Calibri", size=11, bold=True)
        border_bottom = Border(bottom=Side(style='thin'))
        order_ids = self.env['stock.quant'].search([]).filtered(lambda rec: rec.last_count_date and rec.last_count_date.year == eval(self.year))
        table_header = 1

        wrksht = wb.create_sheet(title="Sales Report")
        table_row = table_header + 1
        table_col = 1
        column_loc = 1


        wrksht.cell(row=table_header, column=column_loc).value = 'Item'
        wrksht.cell(row=table_header, column=column_loc).font = bold_text
        wrksht.cell(row=table_header, column=column_loc).alignment = table_header_alignment
        wrksht.cell(row=table_header, column=column_loc).border = border_bottom
        column_loc += 1

        wrksht.cell(row=table_header, column=column_loc).value = 'Description'
        wrksht.cell(row=table_header, column=column_loc).font = bold_text
        wrksht.cell(row=table_header, column=column_loc).alignment = table_header_alignment
        wrksht.cell(row=table_header, column=column_loc).border = border_bottom
        column_loc += 1

        wrksht.cell(row=table_header, column=column_loc).value = 'In Stock'
        wrksht.cell(row=table_header, column=column_loc).font = bold_text
        wrksht.cell(row=table_header, column=column_loc).alignment = table_header_alignment
        wrksht.cell(row=table_header, column=column_loc).border = border_bottom
        column_loc += 1

        total_location = self.env['stock.location'].search([('usage','=','internal')])

        for nm in total_location:
            wrksht.cell(row=table_header, column=column_loc).value = f"{nm.name}"
            wrksht.cell(row=table_header, column=column_loc).font = bold_text
            wrksht.cell(row=table_header, column=column_loc).alignment = table_header_alignment
            wrksht.cell(row=table_header, column=column_loc).border = border_bottom
            column_loc += 1

        wrksht.cell(row=table_header, column=column_loc).value = 'Committed'
        wrksht.cell(row=table_header, column=column_loc).font = bold_text
        wrksht.cell(row=table_header, column=column_loc).alignment = table_header_alignment
        wrksht.cell(row=table_header, column=column_loc).border = border_bottom
        column_loc += 1

        wrksht.cell(row=table_header, column=column_loc).value = 'Ordered'
        wrksht.cell(row=table_header, column=column_loc).font = bold_text
        wrksht.cell(row=table_header, column=column_loc).alignment = table_header_alignment
        wrksht.cell(row=table_header, column=column_loc).border = border_bottom
        column_loc += 1

        wrksht.cell(row=table_header, column=column_loc).value = 'Avaliable'
        wrksht.cell(row=table_header, column=column_loc).font = bold_text
        wrksht.cell(row=table_header, column=column_loc).alignment = table_header_alignment
        wrksht.cell(row=table_header, column=column_loc).border = border_bottom
        column_loc += 1

        wrksht.cell(row=table_header, column=column_loc).value = 'Inventory UoM'
        wrksht.cell(row=table_header, column=column_loc).font = bold_text
        wrksht.cell(row=table_header, column=column_loc).alignment = table_header_alignment
        wrksht.cell(row=table_header, column=column_loc).border = border_bottom
        column_loc += 1

        for month in range(1, 13):
            wrksht.cell(row=table_header, column=column_loc).value = f"{calendar.month_name[month]} - {self.year}"
            wrksht.cell(row=table_header, column=column_loc).font = bold_text
            wrksht.cell(row=table_header, column=column_loc).alignment = table_header_alignment
            wrksht.cell(row=table_header, column=column_loc).border = border_bottom
            column_loc += 1
        wrksht.cell(row=table_header, column=column_loc).value = 'Sales History for last 6 Months'
        wrksht.cell(row=table_header, column=column_loc).font = bold_text
        wrksht.cell(row=table_header, column=column_loc).alignment = table_header_alignment
        wrksht.cell(row=table_header, column=column_loc).border = border_bottom

        for order in order_ids:
            table_col = 1
            wrksht.cell(row=table_row, column=table_col).value = order.product_id.default_code or ''
            wrksht.cell(row=table_row, column=table_col).alignment = table_header_alignment
            table_col += 1

            wrksht.cell(row=table_row, column=table_col).value = order.product_id.name or ''
            wrksht.cell(row=table_row, column=table_col).alignment = table_header_alignment
            table_col += 1

            wrksht.cell(row=table_row, column=table_col).value = order.product_id.qty_available or ''
            wrksht.cell(row=table_row, column=table_col).alignment = table_header_alignment
            table_col += 1

            for nm in total_location:
                if order.location_id.name == nm.name:
                    loc_qty = order.product_id.stock_quant_ids.filtered(lambda loc_qty:loc_qty.location_id == order.location_id)
                    wrksht.cell(row=table_row, column=table_col).value = loc_qty.quantity
                    wrksht.cell(row=table_row, column=table_col).alignment = table_header_alignment
                    table_col += 1
                else:
                    wrksht.cell(row=table_row, column=table_col).value = ''
                    wrksht.cell(row=table_row, column=table_col).alignment = table_header_alignment
                    table_col += 1

            wrksht.cell(row=table_row, column=table_col).value = order.product_id.virtual_available or ''
            wrksht.cell(row=table_row, column=table_col).alignment = table_header_alignment
            table_col += 1

            wrksht.cell(row=table_row, column=table_col).value = order.product_id.sales_count or ''
            wrksht.cell(row=table_row, column=table_col).alignment = table_header_alignment
            table_col += 1

            wrksht.cell(row=table_row, column=table_col).value = order.product_id.qty_available or ''
            wrksht.cell(row=table_row, column=table_col).alignment = table_header_alignment
            table_col += 1

            wrksht.cell(row=table_row, column=table_col).value = order.product_id.uom_id.name or ''
            wrksht.cell(row=table_row, column=table_col).alignment = table_header_alignment
            table_col += 1

            sale_lines = self.env['sale.order.line'].search([('order_id.state','=','sale' ), ('product_id','=', order.product_id.id)])
            for month_in in range(1, 13):
                print("month_in", month_in)
                total_sales = sum(sale_lines.filtered(lambda ord:ord.order_id.date_order.month == month_in).mapped('product_uom_qty'))
                wrksht.cell(row=table_row, column=table_col).value = total_sales or ''
                wrksht.cell(row=table_row, column=table_col).alignment = table_header_alignment
                table_col += 1

            current_date = datetime.now()
            # Calculate the date for 6 months ago
            six_months_ago = current_date - timedelta(days=30 * 6)
            total_sales_six_months = sum(sale_lines.filtered(lambda ord: ord.order_id.date_order.month >= six_months_ago.month).mapped('product_uom_qty'))
            # wrksht.cell(row=table_row, column=table_col + month_in).value = total_sales_six_months or ''
            # wrksht.cell(row=table_row, column=table_col + month_in).alignment = table_header_alignment
            wrksht.cell(row=table_row, column=table_col).value = total_sales_six_months or ''
            wrksht.cell(row=table_row, column=table_col).alignment = table_header_alignment


            wrksht.column_dimensions['A'].width = 16
            wrksht.column_dimensions['B'].width = 14
            wrksht.column_dimensions['C'].width = 14
            wrksht.column_dimensions['D'].width = 14
            wrksht.column_dimensions['E'].width = 14
            wrksht.column_dimensions['F'].width = 14
            wrksht.column_dimensions['G'].width = 14
            wrksht.column_dimensions['H'].width = 14
            wrksht.column_dimensions['I'].width = 14
            wrksht.column_dimensions['J'].width = 16
            wrksht.column_dimensions['K'].width = 16
            wrksht.column_dimensions['L'].width = 16
            wrksht.column_dimensions['M'].width = 14
            wrksht.column_dimensions['N'].width = 14
            wrksht.column_dimensions['O'].width = 14
            wrksht.column_dimensions['P'].width = 14
            wrksht.column_dimensions['R'].width = 14
            wrksht.column_dimensions['Q'].width = 16
            wrksht.column_dimensions['R'].width = 14
            wrksht.column_dimensions['S'].width = 14
            wrksht.column_dimensions['T'].width = 14
            wrksht.column_dimensions['U'].width = 14
            wrksht.column_dimensions['V'].width = 14
            wrksht.column_dimensions['W'].width = 14
            wrksht.column_dimensions['X'].width = 17
            wrksht.column_dimensions['Y'].width = 14
            wrksht.column_dimensions['Z'].width = 17
            wrksht.column_dimensions['AA'].width = 17
            wrksht.column_dimensions['AB'].width = 30

            table_row += 1

        wb.remove(wb.active)

        fp = BytesIO()
        wb.save(fp)
        fp.seek(0)
        data = fp.read()
        fp.close()

        wiz_id = self.env['download.report.wizard'].create({'file': base64.b64encode(data)})

        return {
            'type': 'ir.actions.act_url',
            'url': 'web/content/?model=download.report.wizard&download=true&field=file&id=%s&filename=%s.xlsx' % (
                wiz_id.id, f_name),
            'target': 'self',
        }
