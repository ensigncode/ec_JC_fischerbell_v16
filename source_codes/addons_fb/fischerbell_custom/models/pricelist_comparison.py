# -*- coding: utf-8 -*-
# Part of Odoo. See LICENSE file for full copyright and licensing details.

from odoo import _, api, fields, models, modules
from odoo.exceptions import ValidationError
import base64
import io
import openpyxl
import difflib


class ProductPricelistComparison(models.Model):
    _name = 'product.pricelist.comparison'
    _description = 'Product Pricelist Comparison'
    
    def _get_default_file(self):
        file_path = modules.get_module_resource('fischerbell_custom', 'static/src', 'PriceListValues.xlsx')
        return base64.b64encode(open(file_path, 'rb').read())

    active = fields.Boolean(default=True)
    name = fields.Char('Name', required=True)
    pricelist_id = fields.Many2one('product.pricelist', string='Pricelist', required=True)
    demo_filename = fields.Char(default='PriceListValues')
    demo_file = fields.Binary(readonly=True, default=lambda self: self._get_default_file())
    comparison_file = fields.Binary(required=True)
    filename = fields.Char()
    comparison_line_ids = fields.One2many('product.pricelist.comparison.line', 'comparison_id')
    
    @api.constrains('filename')
    def _check_file_type(self):
        for record in self:
            if record.filename and not record.filename.endswith('.xlsx'):
                raise ValidationError("Only .xlsx files are allowed.")
                
    def read_xlsx_data(self):
        def _get_float_price(price):
            if isinstance(price, (int, float)):
                return price
            else:
                return 0
        
        
        file_data = base64.b64decode(self.comparison_file)
        excel_file = io.BytesIO(file_data)
        book = openpyxl.load_workbook(excel_file)
        sheet = book.active
        product_tmpl_obj = self.env['product.template']
        product_var_obj = self.env['product.product']
        
        data_dict = {}
        price_list_lines = []
        
        try:
            headers = [cell.value for cell in sheet[1]]

            for col in range(1, sheet.max_column + 1):
                key = headers[col - 1]
                data_dict[key] = [sheet.cell(row=row, column=col).value for row in range(2, sheet.max_row + 1)]
            
            keys = data_dict.keys()
            product_tmpl_matches = difflib.get_close_matches('Product', keys)
            product_variant_matches = difflib.get_close_matches('Product Variant', keys)
            price_matches = difflib.get_close_matches('Price', keys)
            if not price_matches:
                raise ValidationError(_("There is no column name match with 'Price' !!!"))
            if product_tmpl_matches:
                product_tmpl_match = product_tmpl_matches[0]
                product_variant_match = product_variant_matches[0]
                price_match = price_matches[0]
                product_tmpl = data_dict[product_tmpl_match]
                product_var = data_dict[product_variant_match]
                prices = data_dict[price_match]
                
                product_variant = False
                for product in product_tmpl:
                    index = product_tmpl.index(product)
                    product_tmpl_id = product_tmpl_obj.search([('name', 'ilike', product)], limit=1)
                    pricelist_item = self.pricelist_id.item_ids.filtered(lambda line: line.product_tmpl_id == product_tmpl_id)
                    # pricelist_products = pricelist_item.mapped('product_id')
                    # if product_variant_matches:
                    #     product_variant = product_var_obj.search([('display_name', 'ilike', product_var[index])], limit=1)
                    old_price = pricelist_item[0].fixed_price if pricelist_item else 0
                    new_price = _get_float_price(prices[index])
                    difference = old_price - new_price
                    if product_tmpl_id:
                        values = {
                            'product_tmpl_id': product_tmpl_id.id,
                            'product_id': product_variant.id if product_variant else False,
                            'pricelist_item': pricelist_item[0].id if pricelist_item else False,
                            'old_price': old_price,
                            'new_price': new_price,
                            'difference': difference,
                        }
                        price_list_lines.append((0, 0, values))
                self.write({'comparison_line_ids': price_list_lines})
            else:
                raise ValidationError(_("There is no column name match with 'Product Variant' or 'Product' !!!"))
        except Exception as e:
            raise ValidationError(_(f"Exception while reading data from XLSX file : {e}"))


class ProductPricelistComparisonLine(models.Model):
    _name = 'product.pricelist.comparison.line'
    _description = 'Product Pricelist Comparison Line'
    _rec_name = 'comparison_id'

    comparison_id = fields.Many2one(comodel_name='product.pricelist.comparison',
                                    ondelete='cascade',
                                    required=True)
    active = fields.Boolean(related='comparison_id.active', store=True)
    product_tmpl_id = fields.Many2one(
        comodel_name='product.template',
        string="Product",
        ondelete='cascade')
    product_id = fields.Many2one(
        comodel_name='product.product',
        string="Product Variant",
        ondelete='cascade')
    operation_type = fields.Selection(selection=[
        ('approve', 'Approve'),
        ('reject', 'Reject')])
    pricelist_item = fields.Many2one('product.pricelist.item')
    old_price = fields.Float(string="Old Price")
    new_price = fields.Float(string="New Price")
    difference = fields.Float(string="Difference")

    def action_apply(self):
        if self.pricelist_item:
            self.pricelist_item.write({'fixed_price': self.new_price})
            self.operation_type = 'approve'
        else:
            list_item = self.env['product.pricelist.item']
            self.pricelist_item = list_item.create({
                'base': 'list_price',
                'applied_on': '1_product',
                'compute_price': 'fixed',
                'product_tmpl_id': self.product_tmpl_id.id,
                'pricelist_id': self.comparison_id.pricelist_id.id,
                'product_id': self.product_id.id,
                'fixed_price': self.new_price,
            })
            self.operation_type = 'approve'

    def action_reject(self):
        self.operation_type = 'reject'
